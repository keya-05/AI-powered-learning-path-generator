import uuid
from dataclasses import dataclass
from pathlib import Path

import fitz  # PyMuPDF
from pptx import Presentation
from sqlalchemy.orm import Session

from app.db.vector_store import get_chunk_collection
from app.models.resource import Resource


@dataclass
class ParentBlock:
    unit_number: int
    unit_title: str
    raw_text: str
    page_or_slide_no: int


def parse_resource(file_path: str, file_type: str) -> list[ParentBlock]:
    if file_type == "pdf":
        return _parse_pdf(file_path)
    if file_type == "pptx":
        return _parse_pptx(file_path)
    if file_type == "docx":
        return _parse_docx(file_path)
    raise ValueError(f"Unsupported file_type for parsing: {file_type}")


def _parse_pdf(file_path: str) -> list[ParentBlock]:
    blocks: list[ParentBlock] = []
    with fitz.open(file_path) as doc:
        for page_index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if not text:
                continue
            blocks.append(
                ParentBlock(
                    unit_number=_infer_unit_number(text, fallback=1),
                    unit_title=_infer_unit_title(text, fallback=f"Page {page_index}"),
                    raw_text=text,
                    page_or_slide_no=page_index,
                )
            )
    return blocks


def _parse_pptx(file_path: str) -> list[ParentBlock]:
    blocks: list[ParentBlock] = []
    presentation = Presentation(file_path)
    for slide_index, slide in enumerate(presentation.slides, start=1):
        text_fragments = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if not text_fragments:
            continue
        combined_text = "\n".join(text_fragments)
        blocks.append(
            ParentBlock(
                unit_number=_infer_unit_number(combined_text, fallback=1),
                unit_title=_infer_unit_title(combined_text, fallback=f"Slide {slide_index}"),
                raw_text=combined_text,
                page_or_slide_no=slide_index,
            )
        )
    return blocks


def _parse_docx(file_path: str) -> list[ParentBlock]:
    # DOCX/mixed formats route through `unstructured` in production; kept as
    # a narrow adapter so ingestion callers don't need to branch on it.
    from unstructured.partition.docx import partition_docx

    elements = partition_docx(filename=file_path)
    blocks: list[ParentBlock] = []
    for index, element in enumerate(elements, start=1):
        text = str(element).strip()
        if not text:
            continue
        blocks.append(
            ParentBlock(
                unit_number=_infer_unit_number(text, fallback=1),
                unit_title=_infer_unit_title(text, fallback=f"Section {index}"),
                raw_text=text,
                page_or_slide_no=index,
            )
        )
    return blocks


def _infer_unit_number(text: str, fallback: int) -> int:
    lowered = text.lower()
    if "unit" in lowered:
        for token in lowered.split():
            if token.isdigit():
                return int(token)
    return fallback


def _infer_unit_title(text: str, fallback: str) -> str:
    first_line = text.splitlines()[0].strip() if text.splitlines() else ""
    return first_line[:150] or fallback


def ingest_resource(db: Session, resource: Resource) -> int:
    """
    Runs parsing -> chunking -> vector indexing for a single resource, marking
    its processing_status along the way. Returns the number of chunks indexed.
    Concept extraction into Neo4j (the graph branch from ARCHITECTURE.md
    section 2.2) is dispatched separately by app/services/graph_indexing.py
    once concept extraction is wired to an LLM provider.
    """
    resource.processing_status = "processing"
    db.commit()

    try:
        parent_blocks = parse_resource(resource.storage_url, resource.file_type)
        chunks = _chunk_blocks(resource, parent_blocks)
        if chunks:
            _index_chunks(chunks)
        resource.processing_status = "indexed"
        db.commit()
        return len(chunks)
    except Exception:
        resource.processing_status = "failed"
        db.commit()
        raise


def _chunk_blocks(
    resource: Resource, parent_blocks: list[ParentBlock], target_tokens: int = 400, overlap_tokens: int = 60
) -> list[dict]:
    chunks: list[dict] = []
    step = max(target_tokens - overlap_tokens, 1)
    for block in parent_blocks:
        tokens = block.raw_text.split()
        for start in range(0, len(tokens), step):
            window = tokens[start : start + target_tokens]
            if not window:
                break
            chunk_id = _deterministic_chunk_id(str(resource.id), block.page_or_slide_no, start)
            chunks.append(
                {
                    "chunk_id": chunk_id,
                    "text": " ".join(window),
                    "course_id": str(resource.course_id),
                    "unit_id": f"{resource.course_id}:{block.unit_number}",
                    "resource_type": resource.resource_type,
                    "resource_id": str(resource.id),
                }
            )
    return chunks


def _deterministic_chunk_id(resource_id: str, block_offset: int, span_start: int) -> str:
    namespace = uuid.UUID("6f9619ff-8b86-d011-b42d-00cf4fc964ff")
    return str(uuid.uuid5(namespace, f"{resource_id}:{block_offset}:{span_start}"))


def _index_chunks(chunks: list[dict]) -> None:
    collection = get_chunk_collection()
    collection.upsert(
        ids=[c["chunk_id"] for c in chunks],
        documents=[c["text"] for c in chunks],
        metadatas=[
            {
                "course_id": c["course_id"],
                "unit_id": c["unit_id"],
                "resource_type": c["resource_type"],
                "resource_id": c["resource_id"],
            }
            for c in chunks
        ],
    )
